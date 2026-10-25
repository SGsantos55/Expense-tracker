"""Tests for Step 11: Export to Excel.

These tests describe what the /expenses/export feature SHOULD do, based on the
spec in .claude/specs/11-export-to-excel.md — not on its implementation.

The route is GET /expenses/export and streams an .xlsx workbook of the current
user's expenses, honouring optional date_from / date_to query params.
"""
import io
import time

import pytest
from openpyxl import load_workbook


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXPORT_URL = "/expenses/export"
EXPECTED_HEADERS = ["Date", "Category", "Description", "Amount"]

# Seed data (see database/db.py seed_db): 8 expenses for demo@spendly.com.
SEED_EXPENSE_COUNT = 8
SEED_TOTAL = 363.98  # 45.50+25+120+35+15.99+89.99+22.50+10


# ------------------------------------------------------------------ #
# Fixtures                                                            #
# ------------------------------------------------------------------ #

@pytest.fixture
def app():
    from app import app as flask_app
    flask_app.config.update({
        "TESTING": True,
        "SECRET_KEY": "test-secret",
    })
    return flask_app


@pytest.fixture
def client(app):
    return app.test_client()


@pytest.fixture
def seed_user_id(app):
    """The demo user ID from the seed data."""
    from database.db import get_db
    conn = get_db()
    user = conn.execute(
        "SELECT id FROM users WHERE email = %s", ("demo@spendly.com",)
    ).fetchone()
    conn.close()
    return user["id"] if user else None


@pytest.fixture
def auth_client(client, seed_user_id):
    """A test client logged in as the demo (seed) user."""
    with client.session_transaction() as sess:
        sess["user_id"] = seed_user_id
    return client


@pytest.fixture
def other_user_id(app):
    """Create a second user who owns a distinctly identifiable expense."""
    from database.db import get_db
    conn = get_db()
    email = f"other{int(time.time() * 1000)}@test.com"
    conn.execute(
        "INSERT INTO users (name, email, password_hash) VALUES (%s, %s, %s)",
        ("Other User", email, "dummyhash"),
    )
    conn.commit()
    uid = conn.execute("SELECT lastval() AS id").fetchone()["id"]
    # A uniquely-named expense that must never leak into another user's export.
    conn.execute(
        "INSERT INTO expenses (user_id, amount, category, date, description) "
        "VALUES (%s, %s, %s, %s, %s)",
        (uid, 999.99, "Shopping", "2026-05-20", "OTHER_USER_SECRET_EXPENSE"),
    )
    conn.commit()
    conn.close()
    return uid


# ------------------------------------------------------------------ #
# Helpers                                                             #
# ------------------------------------------------------------------ #

def _load_ws(response):
    """Parse an xlsx HTTP response into the active openpyxl worksheet."""
    wb = load_workbook(io.BytesIO(response.data))
    return wb.active


def _data_rows(ws):
    """Return all rows (as value lists) of the worksheet."""
    return list(ws.iter_rows(values_only=True))


# ------------------------------------------------------------------ #
# Auth guard                                                         #
# ------------------------------------------------------------------ #

class TestExportAuthGuard:
    def test_unauthenticated_redirects_to_login(self, client):
        response = client.get(EXPORT_URL)
        assert response.status_code == 302, "Logged-out export should redirect"
        assert "/login" in response.location, "Should redirect to /login"

    def test_unauthenticated_does_not_return_xlsx(self, client):
        response = client.get(EXPORT_URL)
        assert response.mimetype != XLSX_MIME, "Logged-out user must not get a workbook"


# ------------------------------------------------------------------ #
# Happy path: HTTP semantics                                         #
# ------------------------------------------------------------------ #

class TestExportHttpResponse:
    def test_authenticated_returns_200(self, auth_client):
        response = auth_client.get(EXPORT_URL)
        assert response.status_code == 200, "Logged-in export should return 200"

    def test_content_type_is_xlsx(self, auth_client):
        response = auth_client.get(EXPORT_URL)
        assert response.mimetype == XLSX_MIME, (
            f"Expected xlsx MIME type, got {response.mimetype!r}"
        )

    def test_content_disposition_is_attachment(self, auth_client):
        response = auth_client.get(EXPORT_URL)
        disposition = response.headers.get("Content-Disposition", "")
        assert "attachment" in disposition.lower(), (
            "Export should be served as a file attachment"
        )

    def test_content_disposition_has_xlsx_filename(self, auth_client):
        response = auth_client.get(EXPORT_URL)
        disposition = response.headers.get("Content-Disposition", "")
        assert ".xlsx" in disposition.lower(), (
            "Attachment filename should end in .xlsx"
        )

    def test_body_is_parseable_as_xlsx(self, auth_client):
        response = auth_client.get(EXPORT_URL)
        # Should not raise — valid xlsx opens in Excel/LibreOffice.
        ws = _load_ws(response)
        assert ws is not None, "Response body should be a valid xlsx workbook"


# ------------------------------------------------------------------ #
# Workbook contents                                                  #
# ------------------------------------------------------------------ #

class TestWorkbookContents:
    def test_header_row_present(self, auth_client):
        ws = _load_ws(auth_client.get(EXPORT_URL))
        header = [cell.value for cell in ws[1]]
        assert header[:4] == EXPECTED_HEADERS, (
            f"First row should be the header {EXPECTED_HEADERS}, got {header[:4]}"
        )

    def test_one_row_per_expense(self, auth_client):
        ws = _load_ws(auth_client.get(EXPORT_URL))
        rows = _data_rows(ws)
        # header + one per expense + total row
        expected = 1 + SEED_EXPENSE_COUNT + 1
        assert len(rows) == expected, (
            f"Expected {expected} rows (header + {SEED_EXPENSE_COUNT} expenses + total), "
            f"got {len(rows)}"
        )

    def test_expense_values_appear(self, auth_client):
        ws = _load_ws(auth_client.get(EXPORT_URL))
        descriptions = [
            r[2] for r in _data_rows(ws)[1:-1]  # skip header and total row
        ]
        assert "Lunch with friends" in descriptions, "Seed expense should be exported"
        assert "Internet bill" in descriptions, "Seed expense should be exported"

    def test_amounts_are_numeric(self, auth_client):
        ws = _load_ws(auth_client.get(EXPORT_URL))
        # Expense rows are everything between header (row 1) and total (last row).
        expense_rows = _data_rows(ws)[1:-1]
        for row in expense_rows:
            amount = row[3]
            assert isinstance(amount, (int, float)), (
                f"Amount column should be numeric, got {amount!r}"
            )

    def test_total_row_sums_amounts(self, auth_client):
        ws = _load_ws(auth_client.get(EXPORT_URL))
        rows = _data_rows(ws)
        expense_rows = rows[1:-1]
        total_row = rows[-1]

        summed = sum(r[3] for r in expense_rows)
        total_value = total_row[3]

        assert isinstance(total_value, (int, float)), "Total should be a number"
        assert total_value == pytest.approx(summed), (
            "Total row should equal the sum of the expense amounts"
        )
        assert total_value == pytest.approx(SEED_TOTAL), (
            f"Total should be {SEED_TOTAL} for the seed data"
        )


# ------------------------------------------------------------------ #
# Date filtering                                                     #
# ------------------------------------------------------------------ #

class TestExportDateFilter:
    def test_date_range_narrows_results(self, auth_client):
        # 2026-05-02 is the only seed expense on that single day (45.50 Food).
        response = auth_client.get(f"{EXPORT_URL}?date_from=2026-05-02&date_to=2026-05-02")
        ws = _load_ws(response)
        rows = _data_rows(ws)
        expense_rows = rows[1:-1]
        assert len(expense_rows) == 1, "Single-day range should export exactly 1 expense"
        assert expense_rows[0][3] == pytest.approx(45.50), "Wrong expense exported"
        assert rows[-1][3] == pytest.approx(45.50), "Total should match the filtered row"

    def test_date_range_excludes_out_of_range_expenses(self, auth_client):
        response = auth_client.get(f"{EXPORT_URL}?date_from=2026-05-01&date_to=2026-05-05")
        ws = _load_ws(response)
        expense_rows = _data_rows(ws)[1:-1]
        # In range: 2026-05-02 (45.50), 05-04 (25), 05-05 (120) => 3 expenses.
        assert len(expense_rows) == 3, "Range 05-01..05-05 should export 3 expenses"
        total = sum(r[3] for r in expense_rows)
        assert total == pytest.approx(190.50), "Filtered total should be 190.50"

    def test_reversed_dates_fall_back_to_all(self, auth_client):
        response = auth_client.get(f"{EXPORT_URL}?date_from=2026-05-15&date_to=2026-05-01")
        assert response.status_code == 200
        expense_rows = _data_rows(_load_ws(response))[1:-1]
        assert len(expense_rows) == SEED_EXPENSE_COUNT, (
            "Reversed dates should fall back to exporting all expenses"
        )

    def test_invalid_date_format_falls_back_to_all(self, auth_client):
        response = auth_client.get(f"{EXPORT_URL}?date_from=not-a-date&date_to=2026-05-05")
        assert response.status_code == 200
        assert response.mimetype == XLSX_MIME, "Invalid dates should still produce a workbook"
        expense_rows = _data_rows(_load_ws(response))[1:-1]
        assert len(expense_rows) == SEED_EXPENSE_COUNT, (
            "Malformed date should fall back to exporting all expenses"
        )

    def test_both_dates_malformed_falls_back_to_all(self, auth_client):
        response = auth_client.get(f"{EXPORT_URL}?date_from=abc&date_to=xyz")
        assert response.status_code == 200
        expense_rows = _data_rows(_load_ws(response))[1:-1]
        assert len(expense_rows) == SEED_EXPENSE_COUNT, (
            "Both malformed dates should fall back to exporting all expenses"
        )


# ------------------------------------------------------------------ #
# Ownership scoping                                                  #
# ------------------------------------------------------------------ #

class TestExportOwnership:
    def test_export_excludes_other_users_expenses(self, auth_client, other_user_id):
        """The demo user's export must never contain another user's expense."""
        response = auth_client.get(EXPORT_URL)
        ws = _load_ws(response)
        descriptions = [r[2] for r in _data_rows(ws)[1:-1]]
        assert "OTHER_USER_SECRET_EXPENSE" not in descriptions, (
            "Export leaked another user's expense"
        )

    def test_other_user_only_sees_their_own_expense(self, client, other_user_id):
        """A different user's export contains only their own rows."""
        with client.session_transaction() as sess:
            sess["user_id"] = other_user_id

        response = client.get(EXPORT_URL)
        assert response.status_code == 200
        ws = _load_ws(response)
        expense_rows = _data_rows(ws)[1:-1]
        descriptions = [r[2] for r in expense_rows]
        assert "OTHER_USER_SECRET_EXPENSE" in descriptions, "Own expense should export"
        assert "Lunch with friends" not in descriptions, (
            "Other user's export must not contain the demo user's expenses"
        )
