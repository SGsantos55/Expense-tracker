"""Excel workbook generation for expense exports — no Flask imports."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font


HEADERS = ["Date", "Category", "Description", "Amount"]


def build_expense_workbook(rows):
    """Build an .xlsx workbook of expenses and return it as bytes.

    Args:
        rows (list[dict]): Expenses with date, category, description, amount.

    Returns:
        bytes: The generated .xlsx workbook.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Expenses"

    header_font = Font(bold=True)
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = header_font

    total = 0.0
    for row in rows:
        amount = row["amount"] or 0.0
        total += amount
        ws.append([
            row["date"],
            row["category"],
            row["description"] or "",
            amount,
        ])

    total_row_index = ws.max_row + 1
    ws.cell(row=total_row_index, column=3, value="Total").font = header_font
    ws.cell(row=total_row_index, column=4, value=total).font = header_font

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
